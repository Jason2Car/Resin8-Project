"""
Joins Stage 3's per-line supplier data with Stage 4's quote results by
Source Row, groups into candidate sets (currently size 1 per the known
upstream limitation), scores them, and writes the comparison out in full -
this sheet IS the "supporting supplier evaluation record" the brief asks
for, not just an intermediate file.
"""
import sys
import os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # bom_pipeline/

from openpyxl import load_workbook, Workbook
from selection import build_comparison, DEFAULT_WEIGHTS

COMPARISON_COLUMNS = ["Source Row", "Manufacturer", "Vendor Status", "ISO 9001", "Match Confidence",
                       "Unit Price", "Lead Time (days)", "Score", "Selected"]


def run(sourced_path: str, quotes_path: str, output_path: str = "comparison_bom.xlsx", weights: dict = None):
    sourced_wb = load_workbook(sourced_path, data_only=True)
    sourced_rows = {r[5]: {  # Source Row is column index 5
        "manufacturer": r[9], "vendor_status": r[11], "iso9001": r[12],
        "match_confidence": r[7],
    } for r in sourced_wb["Sourced"].iter_rows(min_row=2, values_only=True)}

    # quotes_wb's Review Needed already carries sourced_wb's forward (Stage 4
    # merges Stage 3's queue in before writing) - pulling from both here would
    # duplicate every entry.
    review_rows = []

    quotes_wb = load_workbook(quotes_path, data_only=True)
    quotes_by_row = {}
    for ref_id, source_row, unit_price, extended_price, currency, lead_time, ship_date, validity, notes \
            in quotes_wb["Quotes Received"].iter_rows(min_row=2, values_only=True):
        quotes_by_row.setdefault(source_row, []).append({
            "unit_price": unit_price, "lead_time_days": lead_time, "ship_date": ship_date,
            "quote_validity": validity, "notes": notes,
        })
    if "Review Needed" in quotes_wb.sheetnames:
        review_rows += [{"Source Row": r[0], "Issue": r[1]}
                         for r in quotes_wb["Review Needed"].iter_rows(min_row=2, values_only=True) if r[0] is not None]

    comparison_rows = []
    for source_row, base in sourced_rows.items():
        candidates = quotes_by_row.get(source_row, [])
        if not candidates:
            continue  # no quote came back for this line at all - already flagged upstream, nothing to compare

        merged = [{**base, **c} for c in candidates]
        scored = build_comparison(merged, weights)

        for i, c in enumerate(scored):
            comparison_rows.append({
                "Source Row": source_row, "Manufacturer": c["manufacturer"], "Vendor Status": c["vendor_status"],
                "ISO 9001": c["iso9001"], "Match Confidence": c["match_confidence"],
                "Unit Price": c.get("unit_price"), "Lead Time (days)": c.get("lead_time_days"),
                "Score": c["score"], "Selected": i == 0,
            })
        if len(candidates) == 1:
            review_rows.append({"Source Row": source_row,
                                 "Issue": "Only one candidate quote available for this line - "
                                          "nothing to compare against yet (see Stage 2/3 known limitation)."})

    _write_output(comparison_rows, review_rows, output_path, weights or DEFAULT_WEIGHTS)
    return comparison_rows, review_rows


def _write_output(comparison_rows, review_rows, output_path, weights_used):
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    ws.append([f"Weights used: {weights_used}"])
    ws.append(COMPARISON_COLUMNS)
    for r in comparison_rows:
        ws.append([r[c] for c in COMPARISON_COLUMNS])

    ws2 = wb.create_sheet("Review Needed")
    ws2.append(["Source Row", "Issue"])
    for r in review_rows:
        ws2.append([r["Source Row"], r["Issue"]])

    wb.save(output_path)


if __name__ == "__main__":
    rows, review = run("../sourced_bom.xlsx", "../quotes_received.xlsx")
    print(f"{len(rows)} candidate row(s) compared across all lines, {len(review)} flagged for review")
