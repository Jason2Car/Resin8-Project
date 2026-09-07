"""
Channel comes from Stage 3 now (Contact Channel / Contact Email / Portal
URL columns), not a Stage-4-local mock - fixing the gap where Stage 4 had
no way to know how to reach a supplier.

Three paths per line, based on Vendor Status + Contact Channel:
  EMAIL   -> draft_and_send() builds and sends an RFQ, check_replies()
             matches inbox replies back by reference ID
  PORTAL  -> check_portals() logs in and pulls quote text directly - no RFQ
             email step, since portal suppliers with existing accounts
             often just have live pricing, not a request/reply cycle
  (none)  -> flagged for manual outreach, same as before

Portal credentials are read from env vars here for simplicity
(PORTAL_USERNAME/PORTAL_PASSWORD) - a real system needs per-supplier
credentials from a secrets store, not one shared login, since a single
generic account won't work across different suppliers' portals anyway.
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # bom_pipeline/

from openpyxl import load_workbook, Workbook
from rfq_builder import build_rfq
import email_client
import portal_navigator
from llm import extract_quote_fields

RFQ_COLUMNS = ["Source Row", "Name", "Manufacturer", "Channel", "Reference ID", "To", "Sent", "Note"]
QUOTE_COLUMNS = ["Reference ID", "Source Row", "Unit Price", "Extended Price", "Currency",
                  "Lead Time (days)", "Ship Date", "Quote Validity", "Notes"]


def draft_and_send(sourced_path: str, output_path: str = "rfqs_sent.xlsx"):
    wb_in = load_workbook(sourced_path, data_only=True)
    ws = wb_in["Sourced"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    review_rows = []
    if "Review Needed" in wb_in.sheetnames:
        review_rows = [{"Source Row": r[0], "Issue": r[1]}
                        for r in wb_in["Review Needed"].iter_rows(min_row=2, values_only=True) if r[0] is not None]

    rfqs, portal_lines = [], []
    for (name, unit_id, specs, qty, reasoning, source_row, category, match_confidence,
         resolved_unit_id, manufacturer, catalog_description, vendor_status, iso9001,
         channel, contact_email, portal_url) in rows:

        if vendor_status == "UNSOURCED":
            continue
        if vendor_status == "UNEXPLORED_SOURCE":
            # No contact/login on file at all - Stage 3 already flagged this
            # for manual outreach. Nothing for Stage 4 to dispatch to.
            continue
        if category == "ELECTRONIC":
            review_rows.append({"Source Row": source_row,
                                 "Issue": "Distributor part - pricing should come from a Stage 2 catalog "
                                          "query, not an RFQ. Skipped here."})
            continue

        if channel == "PORTAL":
            portal_lines.append({"source_row": source_row, "name": name, "manufacturer": manufacturer,
                                  "portal_url": portal_url})
            continue

        if channel != "EMAIL" or not contact_email:
            review_rows.append({"Source Row": source_row,
                                 "Issue": f"No usable contact/channel for '{manufacturer}' - needs manual outreach."})
            continue

        rfq = build_rfq(name, resolved_unit_id or unit_id, specs, qty, source_row)
        sent = email_client.send_email(contact_email, rfq["subject"], rfq["body"])
        rfqs.append({"Source Row": source_row, "Name": name, "Manufacturer": manufacturer, "Channel": "EMAIL",
                     "Reference ID": rfq["reference_id"], "To": contact_email, "Sent": sent,
                     "Note": "" if sent else "No SMTP configured - RFQ drafted but not sent."})
        if not sent:
            review_rows.append({"Source Row": source_row,
                                 "Issue": f"RFQ to '{manufacturer}' drafted but not sent (no SMTP access)."})

    _write_rfqs(rfqs, review_rows, output_path)
    return rfqs, portal_lines, review_rows


def check_replies(rfqs: list, use_mock: bool = False, output_path: str = "quotes_received.xlsx"):
    reference_ids = [r["Reference ID"] for r in rfqs] if use_mock else [r["Reference ID"] for r in rfqs if r["Sent"]]
    if use_mock:
        import mock_email_backend
        replies = {rid: v for rid, v in mock_email_backend.MOCK_REPLIES.items() if rid in reference_ids}
    else:
        replies = email_client.fetch_replies(reference_ids)

    ref_to_row = {r["Reference ID"]: r["Source Row"] for r in rfqs}
    extracted = extract_quote_fields([{"reference_id": rid, "body_text": d["body"]} for rid, d in replies.items()])
    return _build_quote_rows(replies, extracted, ref_to_row, source="email")


def check_portals(portal_lines: list, use_mock: bool = False):
    """Returns (quotes, review_rows) - separate from check_replies since
    portal quotes aren't tied to an RFQ reference ID, just a source row."""
    review_rows = []
    page_texts = {}

    for line in portal_lines:
        if use_mock:
            import mock_portal_backend
            text = mock_portal_backend.MOCK_PORTAL_PAGES.get(line["manufacturer"])
            result = {"text": text, "stage_reached": "mock", "failure_detail": "",
                      "page_title": None, "page_url": None}
        else:
            result = portal_navigator.fetch_portal_quote(
                line["portal_url"],
                username=os.environ.get("PORTAL_USERNAME"),
                password=os.environ.get("PORTAL_PASSWORD"),
            )

        if result["text"] is None:
            # Full diagnostic captured here on purpose - stage_reached and
            # page_title/url are what let failures be grouped and spotted
            # as a pattern later ("this platform's login page always fails
            # at login_fields_not_found"), not just logged as one-off misses.
            review_rows.append({
                "Source Row": line["source_row"],
                "Issue": f"Could not retrieve a quote from {line['manufacturer']}'s portal "
                         f"({line['portal_url']}) - stopped at '{result['stage_reached']}'"
                         + (f" on page '{result['page_title']}' ({result['page_url']})" if result['page_title'] else "")
                         + (f": {result['failure_detail']}" if result['failure_detail'] else "")
                         + " - needs manual login and check.",
            })
        else:
            page_texts[str(line["source_row"])] = result["text"]

    extracted = extract_quote_fields([{"reference_id": rid, "body_text": t} for rid, t in page_texts.items()])
    ref_to_row = {rid: int(rid) for rid in page_texts}
    quotes, extraction_review = _build_quote_rows(
        {rid: {"body": t} for rid, t in page_texts.items()}, extracted, ref_to_row, source="portal")
    return quotes, review_rows + extraction_review


def _build_quote_rows(replies, extracted, ref_to_row, source):
    quotes, review_rows = [], []
    for rid, data in replies.items():
        source_row = ref_to_row.get(rid)
        fields = extracted.get(rid)
        if fields is None:
            review_rows.append({"Source Row": source_row,
                                 "Issue": f"{source.capitalize()} content received for {rid} but quote fields "
                                          f"could not be extracted (no LLM access) - needs manual read: {data['body'][:200]}"})
            quotes.append({"Reference ID": rid, "Source Row": source_row, "Unit Price": None,
                           "Extended Price": None, "Currency": None, "Lead Time (days)": None,
                           "Ship Date": None, "Quote Validity": None, "Notes": "extraction failed - see review queue"})
        else:
            quotes.append({"Reference ID": rid, "Source Row": source_row,
                           "Unit Price": fields.get("unit_price"), "Extended Price": fields.get("extended_price"),
                           "Currency": fields.get("currency"), "Lead Time (days)": fields.get("lead_time_days"),
                           "Ship Date": fields.get("ship_date"), "Quote Validity": fields.get("quote_validity"),
                           "Notes": fields.get("notes")})
    return quotes, review_rows


def _write_rfqs(rfqs, review_rows, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "RFQs Sent"
    ws.append(RFQ_COLUMNS)
    for r in rfqs:
        ws.append([r[c] for c in RFQ_COLUMNS])
    ws2 = wb.create_sheet("Review Needed")
    ws2.append(["Source Row", "Issue"])
    for r in review_rows:
        ws2.append([r["Source Row"], r["Issue"]])
    wb.save(output_path)


def write_quotes(quotes, review_rows, output_path="quotes_received.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Quotes Received"
    ws.append(QUOTE_COLUMNS)
    for r in quotes:
        ws.append([r[c] for c in QUOTE_COLUMNS])
    ws2 = wb.create_sheet("Review Needed")
    ws2.append(["Source Row", "Issue"])
    for r in review_rows:
        ws2.append([r["Source Row"], r["Issue"]])
    wb.save(output_path)


if __name__ == "__main__":
    rfqs, portal_lines, review1 = draft_and_send("../sourced_bom.xlsx")
    print(f"{len(rfqs)} email RFQ(s) drafted, {len(portal_lines)} portal line(s) to check")

    email_quotes, review2 = check_replies(rfqs, use_mock=True)
    portal_quotes, review3 = check_portals(portal_lines, use_mock=True)

    all_quotes = email_quotes + portal_quotes
    all_review = review1 + review2 + review3
    write_quotes(all_quotes, all_review)
    print(f"{len(email_quotes)} email quote(s), {len(portal_quotes)} portal quote(s), {len(all_review)} flagged for review")
