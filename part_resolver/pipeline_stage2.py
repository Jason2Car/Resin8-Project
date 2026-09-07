"""
Reads Stage 1's output, resolves every line against a catalog, writes
resolved_bom.xlsx. The Review Needed queue is carried forward and appended
to, not replaced.

Two passes, not one, so the LLM plausibility check can be batched:
  pass 1 - find_candidate() per line (local, cheap)
  pass 2 - every exact_id hit goes into ONE llm.check_part_plausibility()
           call, then confidence is finalized using that result

Pass unmocked_backends... use_mock=True to see this run against the sample
data with canned catalog responses standing in for Nexar/a real search
backend, and set an ANTHROPIC_API_KEY to see the plausibility check run for
real instead of failing closed.
"""
import sys
import os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # bom_pipeline/ for llm.py

from openpyxl import load_workbook, Workbook
from resolve import find_candidate
from llm import check_part_plausibility

RESOLVED_COLUMNS = ["Name", "Unit ID", "Specifications/Size", "Quantity", "Reasoning",
                    "Source Row", "Category", "Match Confidence", "Resolved Unit ID",
                    "Manufacturer", "Catalog Description"]


def run(input_path: str, output_path: str = "resolved_bom.xlsx", use_mock: bool = False):
    if use_mock:
        import mock_catalog as electronic_backend
        import mock_catalog as mechanical_backend
    else:
        import nexar_client as electronic_backend
        import mechanical_search as mechanical_backend

    wb_in = load_workbook(input_path, data_only=True)
    std_ws = wb_in["Standardized"]
    rows = list(std_ws.iter_rows(min_row=2, values_only=True))

    review_rows = []
    if "Review Needed" in wb_in.sheetnames:
        review_rows = [{"Source Row": r[0], "Issue": r[1]}
                        for r in wb_in["Review Needed"].iter_rows(min_row=2, values_only=True) if r[0] is not None]

    # Pass 1: find a candidate for every line.
    lines = []
    for name, unit_id, specs, qty, reasoning, source_row in rows:
        name, unit_id, specs, reasoning = (v or "" for v in (name, unit_id, specs, reasoning))
        candidate = find_candidate(name, unit_id, specs, electronic_backend, mechanical_backend)
        lines.append({"name": name, "unit_id": unit_id, "specs": specs, "qty": qty,
                       "reasoning": reasoning, "source_row": source_row, **candidate})

    # Pass 2: batch-check plausibility for every exact_id hit in one call.
    exact_id_lines = [l for l in lines if l["source"] == "exact_id"]
    plausibility = check_part_plausibility([
        {"source_row": l["source_row"], "bom_name": l["name"], "catalog_description": l["catalog_description"]}
        for l in exact_id_lines
    ])

    resolved = []
    for l in lines:
        if l["source"] == "exact_id":
            verdict = plausibility.get(l["source_row"])
            if verdict is None:
                # No LLM available - fail closed: don't assume EXACT, flag it.
                confidence, note = "PROBABLE_EQUIVALENT", (
                    "Unit ID found in catalog, but plausibility could not be confirmed "
                    "(no LLM access) - verify this is the right part before treating it as exact.")
            elif verdict["plausible"]:
                confidence, note = "EXACT_MATCH", ""
            else:
                confidence, note = "PROBABLE_EQUIVALENT", (
                    f"Unit ID found, but catalog description doesn't clearly match the BOM name: {verdict['note']}")
        elif l["source"] == "description_search":
            confidence, note = "PROBABLE_EQUIVALENT", (
                "Matched by description only, no manufacturer part number was available/verifiable - "
                "needs engineering sign-off before treating as equivalent.")
        else:
            confidence, note = "NO_MATCH", "No exact or probable match found - route to supplier discovery from scratch."

        resolved.append({
            "Name": l["name"], "Unit ID": l["unit_id"], "Specifications/Size": l["specs"],
            "Quantity": l["qty"], "Reasoning": l["reasoning"], "Source Row": l["source_row"],
            "category": l["category"], "match_confidence": confidence,
            "resolved_unit_id": l["resolved_unit_id"], "manufacturer": l["manufacturer"],
            "catalog_description": l["catalog_description"],
        })
        if note:
            review_rows.append({"Source Row": l["source_row"], "Issue": note})

    _write_output(resolved, review_rows, output_path)
    return resolved, review_rows


def _write_output(resolved, review_rows, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resolved"
    ws.append(RESOLVED_COLUMNS)
    for r in resolved:
        ws.append([r["Name"], r["Unit ID"], r["Specifications/Size"], r["Quantity"], r["Reasoning"],
                    r["Source Row"], r["category"], r["match_confidence"],
                    r["resolved_unit_id"], r["manufacturer"], r["catalog_description"]])

    ws2 = wb.create_sheet("Review Needed")
    ws2.append(["Source Row", "Issue"])
    for r in review_rows:
        ws2.append([r["Source Row"], r["Issue"]])

    wb.save(output_path)


if __name__ == "__main__":
    resolved, review = run("../standardized_bom.xlsx", use_mock=True)
    counts = {}
    for r in resolved:
        counts[r["match_confidence"]] = counts.get(r["match_confidence"], 0) + 1
    print(counts)
    print(f"{len(review)} rows in review queue")
